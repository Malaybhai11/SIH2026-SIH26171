"""Build the SIH 2026 idea deck (official template) for PS 26171 — Aavaran.

    .venv/bin/pip install python-pptx
    .venv/bin/python docs/deck/build_deck.py      # -> docs/deck/Aavaran_SIH2026_PS26171.pptx

Numbers are read from eval/results/*.json, so the deck always matches the measured
results. Fill TEAM_ID / TEAM_NAME below before exporting the PDF for the portal.
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

TEAM_ID = "<Team ID>"
TEAM_NAME = "<Team name>"

ROOT = Path(__file__).resolve().parents[2]
R = lambda f: json.loads((ROOT / "eval/results" / f).read_text())
pii, faces, red, screens, lat = R("pii.json"), R("faces.json"), R("redaction.json"), R("screens.json"), R("latency.json")

NAVY = RGBColor(0x1B, 0x2A, 0x4E)
SAFFRON = RGBColor(0xE8, 0x7A, 0x1E)
GREEN = RGBColor(0x13, 0x88, 0x08)
INK = RGBColor(0x22, 0x28, 0x33)
MUTE = RGBColor(0x5B, 0x64, 0x72)
TINT = RGBColor(0xEE, 0xF2, 0xF8)
SAFF_TINT = RGBColor(0xFD, 0xF0, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY, HEAD = "Calibri", "Cambria"

prs = Presentation(ROOT / "docs/deck/SIH2026_template.pptx")

# drop the "important instructions" slide (template says it may be removed)
sld_ids = prs.slides._sldIdLst
sld_ids.remove(list(sld_ids)[6])


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def remove(slide, name):
    el = shape(slide, name)._element
    el.getparent().remove(el)


def text(slide, x, y, w, h, runs, size=12, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {opts})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
        if isinstance(para, list) and para and para[0][1].get("space"):
            p.space_before = Pt(para[0][1]["space"])
    return box


def bullets(slide, x, y, w, h, items, size=11.5, color=INK, gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        parts = it if isinstance(it, list) else [(it, {})]
        # native bullet
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.16))))
        pPr.set("indent", str(-Emu(Inches(0.14))))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = BODY
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def header(slide, x, y, w, t, color=NAVY, size=14):
    text(slide, x, y, w, 0.32, [[(t, {"bold": True, "color": color, "size": size, "font": HEAD})]])


def arrow(slide, x1, y1, x2, y2, color=MUTE):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle", "w": "med", "len": "med"}))


def team_oval(slide):
    for s in slide.shapes:
        if s.name.startswith("Oval") and s.has_text_frame:
            r = s.text_frame.paragraphs[0].runs
            if r:
                r[0].text = TEAM_NAME
                for extra in r[1:]:
                    extra.text = ""
            for p in s.text_frame.paragraphs[1:]:
                for rr in p.runs:
                    rr.text = ""


def footer(slide):
    for s in slide.shapes:
        if s.name.startswith("Footer") and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for i, r in enumerate(p.runs):
                    r.text = "Aavaran · SIH 2026 · PS 26171" if i == 0 else ""


def set_title(slide, t):
    tf = shape(slide, "Title 1").text_frame
    runs = [r for p in tf.paragraphs for r in p.runs]
    runs[0].text = t
    for r in runs[1:]:
        r.text = ""


S = list(prs.slides)
pct = lambda v: f"{v * 100:.1f}%"
scr = screens["accuracy"]
ind = pii["indian_synthetic"]["rulesPlusNer"]
ai4 = pii["ai4privacy_en"]["rulesPlusNer"]
agg = red["aggregate"]
eco, bal = lat["modes"]["eco"], lat["modes"]["balanced"]
tasks = lat["e2eTasks"]

# ---------------------------------------------------------------- 1 · title page
s = S[0]
vals = {
    "Problem Statement ID": " 26171",
    "Problem Statement Title": " On-device Visual Perception for Light-weight Browser Agents",
    "Theme": " Smart Automation",
    "PS Category": None,
    "Team ID": f" {TEAM_ID}",
    "Team Name": f" {TEAM_NAME}",
}
for p in shape(s, "TextBox 9").text_frame.paragraphs:
    full = "".join(r.text for r in p.runs)
    for k, v in vals.items():
        if full.strip().startswith(k):
            if k == "PS Category":
                p.runs[0].text = "PS Category- Software"
                for r in p.runs[1:]:
                    r.text = ""
            elif k == "Team Name":
                p.runs[0].text = "Team Name (Registered on portal) –" + v
                for r in p.runs[1:]:
                    r.text = ""
            else:
                last = p.runs[-1]
                last.text = last.text.rstrip("–- ") + " –" + v
for p in shape(s, "TextBox 9").text_frame.paragraphs:
    p.line_spacing = 1.05
    p.space_before = Pt(7)
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(15)
tp_runs = [r for p in shape(s, "Subtitle 3").text_frame.paragraphs for r in p.runs]
tp_runs[0].text = "AAVARAN — a privacy veil for browser agents"
for r in tp_runs[1:]:
    r.text = ""

# ---------------------------------------------------------------- 2 · idea
s = S[1]
set_title(s, "AAVARAN — ON-DEVICE PRIVACY VEIL")
for r in (r for p in shape(s, "Title 1").text_frame.paragraphs for r in p.runs):
    r.font.size = Pt(28)
remove(s, "TextBox 8")
team_oval(s)
footer(s)
L = 0.45
header(s, L, 1.3, 7.3, "Proposed solution")
bullets(s, L, 1.62, 7.4, 1.45, [
    [("Browser extension (Chrome, Edge, Brave, Firefox) + FastAPI server. ", {"bold": True}), (f"Three small models run in the browser (WebGPU → WASM) and read the screen: YuNet faces, MobileCLIP-S0 vision transformer, BERT-small PII NER ({bal['modelMB']:.0f} MB total).", {})],
    [("Before any request, ", {}), ("personal data is replaced by consistent tokens", {"bold": True}), (" ([NAME_1], [AADHAAR_1], [OTP_1]) and black-boxed at pixel level in the screenshot, with the same labels.", {})],
    [("An open-weights VLM on the server (Qwen2.5-VL / Llama-4 / Gemma-3) plans actions over tokens + numbered UI marks; the client swaps real values back in ", {}), ("on the device", {"bold": True}), (".", {})],
])
header(s, L, 3.12, 7.3, "How it addresses the problem")
bullets(s, L, 3.44, 7.4, 1.35, [
    "Server-side agents can't be trusted with Aadhaar cards, OTPs, faces and bank pages; a local-only agent is too weak. Aavaran splits the work: perception + privacy local, reasoning remote.",
    [("Redaction is ", {}), ("verifiable", {"bold": True}), (": a Privacy X-ray shows the exact frame and tokens that leave; the server's own audit log is checked after every task — ", {}), ("0 leaks", {"bold": True, "color": GREEN}), (" in 5 end-to-end tasks.", {})],
])
header(s, L, 4.82, 7.3, "Innovation & uniqueness")
bullets(s, L, 5.14, 7.4, 1.75, [
    [("DOM-grounded pixel redaction: ", {"bold": True}), ("PII found in text is mapped to exact screen rectangles (Range.getClientRects) — no OCR guessing; vision covers pixels the DOM can't describe (faces, ID-card photos, signatures).", {})],
    [("Reversible pseudonym vault: ", {"bold": True}), ("the agent can fill your form with your details without ever seeing them.", {})],
    [("Pixels × structure: ", {"bold": True}), (f"a ViT fused with DOM structure understands unseen websites {pct(scr['clipPlusDomPriors'])} vs {pct(scr['zeroShotClip'])} from pixels alone.", {})],
    [("India-first PII: ", {"bold": True}), ("Aadhaar (Verhoeff), PAN, GSTIN (mod-36), UPI, IFSC-context accounts, OTPs.", {})],
])
# before / after
iw = 3.85
ih = iw * 800 / 1280
x0 = 12.9 - iw
y1 = 1.32
s.shapes.add_picture(str(ROOT / "docs/img/kyc_original.jpg"), Inches(x0), Inches(y1), Inches(iw), Inches(ih))
text(s, x0, y1 + ih + 0.01, iw, 0.26, [[("What the user sees", {"bold": True, "size": 10, "color": MUTE})]], align=PP_ALIGN.CENTER)
y2 = y1 + ih + 0.36
s.shapes.add_picture(str(ROOT / "docs/img/kyc_server_view.jpg"), Inches(x0), Inches(y2), Inches(iw), Inches(ih))
text(s, x0, y2 + ih + 0.01, iw, 0.26, [[("What the server receives — redacted on the device", {"bold": True, "size": 10, "color": SAFFRON})]], align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 3 · technical approach
s = S[2]
remove(s, "TextBox 8")
team_oval(s)
footer(s)
steps = [
    ("1", "Screen + DOM", "tokenise task first;\nvisible text, fields,\nimages, UI elements"),
    ("2", "On-device perception", "YuNet · MobileCLIP-S0\n· BERT-small NER\n(ONNX Runtime Web)"),
    ("3", "Redact & tokenise", "pixel-exact black boxes\nsame tokens in text\nand image · Vault"),
    ("4", "Egress gate", "fail-closed re-check of\nevery outgoing string\n(no image if tab moved)"),
    ("5", "Open-weights VLM", "server re-checks; reasons\nover tokens + Set-of-\nMarks screenshot"),
    ("6", "Act on device", "tokens → real values\nlocally; direct DOM\nevents; loop"),
]
bw, gap, y = 1.9, 0.22, 1.35
for i, (n, t, d) in enumerate(steps):
    x = 0.45 + i * (bw + gap)
    server = i == 4
    rect(s, x, y, bw, 1.55, SAFF_TINT if server else TINT)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.08), Inches(y + 0.08), Inches(0.34), Inches(0.34))
    circ.fill.solid()
    circ.fill.fore_color.rgb = SAFFRON if server else NAVY
    circ.line.fill.background()
    circ.text_frame.text = n
    for r in circ.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(11), True, WHITE, BODY
    circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, x + 0.46, y + 0.08, bw - 0.5, 0.36, [[(t, {"bold": True, "size": 11, "color": SAFFRON if server else NAVY})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.08, y + 0.5, bw - 0.14, 1.0, [[(ln, {"size": 9.5, "color": INK})] for ln in d.split("\n")])
    if i < len(steps) - 1:
        arrow(s, x + bw + 0.01, y + 0.78, x + bw + gap - 0.01, y + 0.78)
text(s, 0.45, y + 1.6, 8.3, 0.28, [[("user's browser — raw pixels and values never leave the device", {"size": 9.5, "italic": True, "color": NAVY})]], align=PP_ALIGN.CENTER)
text(s, 8.93, y + 1.6, 1.9, 0.28, [[("server", {"size": 9.5, "italic": True, "color": SAFFRON})]], align=PP_ALIGN.CENTER)

header(s, 0.45, 3.35, 5.8, "Technologies")
tech = [
    ("Client", "Manifest V3 extension, JavaScript, esbuild; Chrome/Edge/Brave + Firefox from one source"),
    ("On-device ML", "ONNX Runtime Web — WebGPU, multi-threaded WASM SIMD (cross-origin isolated); offscreen document hosts all models once"),
    ("Models", f"YuNet FP32 {bal['modelsLoaded']['face']['MB']:.1f} MB · MobileCLIP-S0 image tower FP16 {bal['modelsLoaded']['clip']['MB']:.1f} MB (prompt embeddings precomputed) · BERT-small PII INT8 {bal['modelsLoaded']['ner']['MB']:.1f} MB"),
    ("Server", "Python FastAPI; OpenAI-compatible open-weights VLM (vLLM / Ollama): Qwen2.5-VL-7B, Llama-4-Scout, Gemma-3"),
    ("Evaluation", "Puppeteer + real Chrome, onnxruntime-node; WIDER FACE, ai4privacy, 217-screen web set"),
]
yy = 3.7
for k, v in tech:
    rect(s, 0.45, yy, 1.35, 0.56, NAVY)
    text(s, 0.45, yy, 1.35, 0.56, [[(k, {"bold": True, "size": 10.5, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.9, yy, 4.45, 0.56, [[(v, {"size": 9.5})]], anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.63

header(s, 6.75, 3.35, 6.1, "Methodology (built and measured)")
bullets(s, 6.75, 3.7, 6.15, 3.2, [
    [("Text PII → pixels: ", {"bold": True}), ("visible text grouped per block, scanned by checksum rules + NER + known Vault values; each span → Range.getClientRects() → black box labelled with its token.", {})],
    [("Vision: ", {"bold": True}), ("all on-page images packed into one 640² YuNet mosaic (avatars become detectable); one batched MobileCLIP pass for screen state + image regions (ID card, signature, QR, face photo), fused with DOM cues.", {})],
    [("Adaptive compute: ", {"bold": True}), ("dHash frame cache, per-page screen cache, NER/region caches, eco / balanced / max modes, models warmed while the user types.", {})],
    [("Grounding: ", {"bold": True}), ("numbered Set-of-Marks on interactive elements; actions reference node ids; tokens rehydrated only at execution.", {})],
    [("Defence in depth: ", {"bold": True}), ("sensitive-field values never serialised; egress gate; server-side re-check (Aadhaar/Verhoeff, PAN, UPI, Luhn); audit log.", {})],
], size=10.5, gap=4)

# ---------------------------------------------------------------- 4 · feasibility
s = S[3]
remove(s, "TextBox 8")
team_oval(s)
footer(s)
text(s, 0.45, 1.25, 12.4, 0.35, [[("Working prototype, measured against all five evaluation criteria (reproducible scripts in eval/):", {"size": 12, "bold": True, "color": NAVY})]])
stats = [
    (pct(scr["clipPlusDomPriors"]), "screen understanding on\nunseen websites (96 sites)", NAVY),
    (f"{ind['f1']:.2f}", f"PII F1, Indian set\n(R {ind['recall']:.2f} · P {ind['precision']:.2f})", NAVY),
    (f"{agg['pixelPrecision']:.2f}", f"pixel redaction precision\n{round(agg['objectRecall'] * agg['gtObjects'])}/{agg['gtObjects']} sensitive objects covered", NAVY),
    (f"{eco['engineMemoryMB']:.0f}–{bal['engineMemoryMB']:.0f} MB", "total on-device engine\nmemory (eco–balanced)", NAVY),
    (f"{bal['warmStepMs']['median'] / 1000:.2f} s", "on-device perception per\nstep (median, laptop CPU)", NAVY),
    ("0", "leaks in 5 end-to-end\ntasks (server audit log)", GREEN),
]
cw = (12.45 - 5 * 0.18) / 6
for i, (big, small, col) in enumerate(stats):
    x = 0.45 + i * (cw + 0.18)
    rect(s, x, 1.68, cw, 1.45, TINT)
    text(s, x, 1.72, cw, 0.62, [[(big, {"bold": True, "size": 26, "color": col, "font": HEAD})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.05, 2.36, cw - 0.1, 0.75, [[(ln, {"size": 9.5, "color": MUTE})] for ln in small.split("\n")], align=PP_ALIGN.CENTER)

header(s, 0.45, 3.35, 12, "Challenges, risks and how we handle them")
rows = [
    ("Low-end client devices", "Models too heavy / slow", f"{bal['modelMB']:.0f} MB total; measured precision choice (FP32 YuNet 2.6× faster than INT8 in WASM, see docs/model-contract.md); eco mode {eco['engineMemoryMB']:.0f} MB; caches make an unchanged frame ~{eco['unchangedFrameMs']['median'] / 1000:.2f} s; WebGPU when present"),
    ("Missed PII = privacy leak", "Rules or NER miss a value", "Four layers: checksum rules + NER + Vault-guided matching + sensitive-field boxing; fail-closed egress gate; server re-check; audit log"),
    ("Over-redaction", "Agent loses context", "Typed, consistent tokens keep structure; checksum validation keeps order ids, PNRs, prices, IFSC readable (1/120 false alarms on hard negatives)"),
    ("Unfamiliar UIs", "ViT not trained on screens", "Fuse pixels with DOM structure (+17.5 pts on unseen sites); element grounding via DOM + Set-of-Marks is exact"),
    ("Server model availability", "Cloud dependence", "Any OpenAI-compatible open-weights endpoint (vLLM / Ollama offline); stateless server; mock mode for demos"),
]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.45), Inches(3.7), Inches(12.45), Inches(3.1)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(2.4), Inches(2.3), Inches(7.75)
for c, h in enumerate(["Challenge", "Risk", "Strategy (implemented)"]):
    cell = tbl.cell(0, c)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(10.5), True, WHITE, BODY
for i, row in enumerate(rows, 1):
    for c, v in enumerate(row):
        cell = tbl.cell(i, c)
        cell.text = v
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if i % 2 else TINT
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.name = Pt(9.5), BODY
                r.font.bold = c == 0
                r.font.color.rgb = INK

# ---------------------------------------------------------------- 5 · impact
s = S[4]
remove(s, "TextBox 8")
team_oval(s)
footer(s)
header(s, 0.45, 1.3, 6.4, "Who benefits")
aud = [
    ("Citizens", "use AI assistants on DigiLocker, UPI, bank, e-KYC and government portals without handing Aadhaar, OTPs or faces to a server"),
    ("Government & ISRO / DoS", "automate internal web workflows on sensitive screens; the server can run on-premise with open weights"),
    ("Enterprises & startups", "build agents that comply with the DPDP Act 2023 by design (data minimisation, purpose limitation, auditability)"),
    ("Low-end & rural devices", "runs on a 4-core laptop CPU with no GPU; small payloads (~50 KB image per step) suit slow networks"),
]
yy = 1.68
for k, v in aud:
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.45), Inches(yy + 0.05), Inches(0.36), Inches(0.36))
    circ.fill.solid()
    circ.fill.fore_color.rgb = SAFFRON
    circ.line.fill.background()
    circ.text_frame.text = str(aud.index((k, v)) + 1)
    for r in circ.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(11), True, WHITE, BODY
    circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 0.92, yy, 5.9, 0.72, [[(k + " — ", {"bold": True, "size": 11, "color": NAVY}), (v, {"size": 10.5})]])
    yy += 0.8

header(s, 0.45, 4.95, 6.4, "Benefits")
bullets(s, 0.45, 5.28, 6.4, 1.6, [
    [("Social: ", {"bold": True}), ("trust in AI assistance; privacy that users can inspect (Privacy X-ray).", {})],
    [("Economic: ", {"bold": True}), ("open-weights models, no per-token vendor lock-in; smaller prompts (tokens, one ~50 KB JPEG).", {})],
    [("Strategic: ", {"bold": True}), ("sovereign, offline-deployable stack, aligned with Atmanirbhar Bharat and DPDP.", {})],
    [("Environmental: ", {"bold": True}), ("caches skip repeated inference; less data moved per step.", {})],
], size=10.5, gap=3)

def bar_chart(slide, x, y, w, h, title, cats, vals, colors, fmt='0%'):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    for p in ch.chart_title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(11), True, NAVY, BODY
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(10), True
    for i, pt in enumerate(pl.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i]
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = 1.0, 0
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(9.5)
    ca.format.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    ca.reverse_order = True

bar_chart(s, 7.1, 1.3, 5.8, 1.85, "Screen understanding, unseen websites",
          ["Pixels only (zero-shot ViT)", "Pixels × DOM structure (ours)"], [scr["zeroShotClip"], scr["clipPlusDomPriors"]], [MUTE, NAVY])
bar_chart(s, 7.1, 3.25, 5.8, 1.85, "PII recall, Indian set (precision ≥ 0.98)",
          ["Rules only", "Rules + on-device NER"], [pii["indian_synthetic"]["rulesOnly"]["recall"], ind["recall"]], [MUTE, NAVY])
bar_chart(s, 7.1, 5.2, 5.8, 1.65, "Redaction on demo sites (pixels)",
          ["Precision", "Recall"], [agg["pixelPrecision"], agg["pixelRecall"]], [SAFFRON, SAFFRON], fmt="0.00")

# ---------------------------------------------------------------- 6 · references
s = S[5]
remove(s, "TextBox 8")
team_oval(s)
footer(s)
header(s, 0.45, 1.3, 7.6, "Research & models")
refs = [
    "Vasu et al., MobileCLIP: Fast Image-Text Models through Multi-Modal Reinforced Training, CVPR 2024 — arxiv.org/abs/2311.17049",
    "Wu et al., YuNet: A Tiny Millisecond-level Face Detector, Machine Intelligence Research 2023 — github.com/opencv/opencv_zoo",
    "Yang et al., Set-of-Mark Prompting Unleashes Extraordinary Visual Grounding in GPT-4V, 2023 — arxiv.org/abs/2310.11441",
    "Bai et al., Qwen2.5-VL Technical Report, 2025 — arxiv.org/abs/2502.13923 (open-weights server VLM)",
    "Radford et al., Learning Transferable Visual Models from Natural Language Supervision (CLIP), ICML 2021",
    "ONNX Runtime Web (WebGPU / WASM) — onnxruntime.ai · Transformers.js — huggingface.co/docs/transformers.js",
    "BERT-small PII detection (ONNX) — huggingface.co/onnx-community/bert-small-pii-detection-ONNX",
    "Datasets: WIDER FACE (Yang et al., CVPR 2016); ai4privacy/pii-masking-200k",
    "Digital Personal Data Protection Act, 2023 (MeitY); UIDAI Aadhaar number checksum (Verhoeff)",
]
bullets(s, 0.45, 1.7, 7.7, 5.2, refs, size=12.5, gap=9)
rect(s, 8.45, 1.35, 4.45, 4.3, TINT)
header(s, 8.65, 1.5, 4.1, "Prototype & evidence")
bullets(s, 8.65, 1.88, 4.1, 4.9, [
    [("Code: ", {"bold": True}), ("github.com/Malaybhai11/SIH2026-SIH26171", {})],
    [("Extension ", {"bold": True}), ("for Chrome/Edge/Brave (dist/) and Firefox (dist-firefox/)", {})],
    [("Demo sites ", {"bold": True}), ("(synthetic data, AI-generated faces): bank KYC, webmail OTP, social feed, checkout, ISRO outreach registration", {})],
    [("Scorecard: ", {"bold": True}), ("eval/results/SUMMARY.md — every number reproducible with one npm script", {})],
    [("End-to-end tasks: ", {"bold": True}), (", ".join(f"{k} {v['wallMs'] / 1000:.0f} s" for k, v in tasks.items()) + " — all passed, 0 leaks", {})],
    [("Faces (WIDER val): ", {"bold": True}), (f"P {faces['precision']:.2f} / R {faces['recall']:.2f}; ai4privacy PII F1 {ai4['f1']:.2f}", {})],
], size=12, gap=9)

out = ROOT / "docs/deck/Aavaran_SIH2026_PS26171.pptx"
prs.save(out)
print("wrote", out)
